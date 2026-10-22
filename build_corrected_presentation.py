"""
build_corrected_presentation.py - Generate updated, verified 12-slide PowerPoint presentation with honest trade-offs and speaker notes.
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_deck(output_path):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    NAVY = RGBColor(24, 43, 73)
    SLATE = RGBColor(52, 73, 94)
    ACCENT_BLUE = RGBColor(41, 128, 185)
    LIGHT_BG = RGBColor(245, 247, 250)
    DARK_TEXT = RGBColor(33, 33, 33)
    WHITE = RGBColor(255, 255, 255)
    GREEN = RGBColor(39, 174, 96)
    RED = RGBColor(192, 57, 43)

    blank_layout = prs.slide_layouts[6]

    def add_header(slide, title_text, category="IBM3201 DATA MINING & PREDICTIVE ANALYTICS"):
        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(1.1))
        tf = header_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p_cat = tf.paragraphs[0]
        p_cat.text = category.upper()
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = ACCENT_BLUE

        p_title = tf.add_paragraph()
        p_title.text = title_text
        p_title.font.size = Pt(24)
        p_title.font.bold = True
        p_title.font.color.rgb = NAVY

    # SLIDE 1: Title Slide
    s1 = prs.slides.add_slide(blank_layout)
    bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = NAVY
    bg1.line.color.rgb = NAVY

    tb1 = s1.shapes.add_textbox(Inches(1.0), Inches(1.5), Inches(11.333), Inches(4.5))
    tf1 = tb1.text_frame
    tf1.word_wrap = True

    p = tf1.paragraphs[0]
    p.text = "IBM3201 / BDS32043 DATA MINING & PREDICTIVE ANALYTICS"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE

    p2 = tf1.add_paragraph()
    p2.text = "Credit Card Fraud Detection Using Isolation Forest\nwith Improved Threshold Tuning"
    p2.font.size = Pt(32)
    p2.font.bold = True
    p2.font.color.rgb = WHITE
    p2.space_before = Pt(16)
    p2.space_after = Pt(20)

    p3 = tf1.add_paragraph()
    p3.text = "Case Study Analysis and Improvement of Anomaly Detection Techniques | Group 5"
    p3.font.size = Pt(16)
    p3.font.color.rgb = RGBColor(189, 195, 199)

    p4 = tf1.add_paragraph()
    p4.text = (
        "Group Members:\n"
        "• Mohan Mekala (I26043589)          • Siddrtha Reddy (I26043588)\n"
        "• Boori Koteswara Rao (I26043120)   • Sindu Sri Chilakanti (I26043217)\n"
        "• Gayathri Sai Tapaswini (I26043735)\n\n"
        "Lecturer: Ms. Harprith Kaur | Faculty of Data Science and Technology (FDSIT)"
    )
    p4.font.size = Pt(13)
    p4.font.color.rgb = WHITE
    p4.space_before = Pt(24)

    s1.notes_slide.notes_text_frame.text = (
        "Good day Ms. Harprith and colleagues. We are Group 5, presenting our project: "
        "'Credit Card Fraud Detection Using Isolation Forest with Improved Threshold Tuning'. "
        "We surveyed 10 peer-reviewed anomaly detection studies, selected the benchmark work of Ounacer et al. (2018), "
        "identified an acute operational problem with default decision thresholds under extreme class imbalance, "
        "and implemented a validation-driven threshold optimization that reduces false-positive alerts by 92.79% while raising precision by 444.5%."
    )

    # SLIDE 2: Background & Problem Statement
    s2 = prs.slides.add_slide(blank_layout)
    add_header(s2, "1. Background & Problem Statement")
    tb2 = s2.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
    tf2 = tb2.text_frame
    tf2.word_wrap = True

    p = tf2.paragraphs[0]
    p.text = "The Real-World Financial Challenge:"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = NAVY

    points = [
        ("Massive Digital Transaction Streams: ", "Modern financial institutions process millions of electronic transactions daily, where fraud causes tens of billions in annual losses."),
        ("Severe Class Imbalance: ", "Fraud constitutes less than 0.2% of all payment activity. Genuine transactions outnumber frauds by over 570 to 1."),
        ("Limitations of Static Rule Engines: ", "Heuristic rule sets are rigid and easily circumvented by dynamic criminal tactics like automated bot testing."),
        ("The Burden of Alert Fatigue: ", "When machine learning anomaly models produce thousands of false alarms, human fraud investigator queues overflow, operational overhead skyrockets, and legitimate cardholders suffer embarrassing transaction declines.")
    ]
    for b, n in points:
        p = tf2.add_paragraph()
        p.space_before = Pt(12)
        r1 = p.add_run()
        r1.text = "• " + b
        r1.font.bold = True
        r1.font.size = Pt(14)
        r1.font.color.rgb = NAVY
        r2 = p.add_run()
        r2.text = n
        r2.font.size = Pt(14)
        r2.font.color.rgb = DARK_TEXT

    s2.notes_slide.notes_text_frame.text = (
        "In electronic banking, fraud detection is a needle-in-a-haystack problem. Because frauds are fewer than 0.2% of transactions, "
        "an algorithm that produces even a 3% false alarm rate will overwhelm human investigation teams with thousands of bogus alerts. "
        "Our core problem statement focuses on calibrating anomaly thresholds to make machine learning models operationally viable."
    )

    # SLIDE 3: Objectives
    s3 = prs.slides.add_slide(blank_layout)
    add_header(s3, "2. Project Objectives")
    tb3 = s3.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
    tf3 = tb3.text_frame
    tf3.word_wrap = True

    p = tf3.paragraphs[0]
    p.text = "Core Research & Engineering Goals:"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = NAVY

    objs = [
        ("1. Literature Review: ", "Investigate 10 real-world anomaly detection case studies across banking, cybersecurity, IoT, healthcare, and aerospace."),
        ("2. Replicate Baseline Approach: ", "Implement the Isolation Forest baseline approach of Ounacer et al. (2018) using the European Credit Card Fraud benchmark."),
        ("3. Identify Operational Flaw: ", "Demonstrate why default offset thresholds (0.500) fail in production by generating unacceptable false alarm volumes."),
        ("4. Develop Principled Improvement: ", "Design a validation-based threshold tuning protocol that optimizes the operational F1-score without test data leakage."),
        ("5. Empirical Evaluation: ", "Evaluate both approaches on identical test records and honestly analyze the precision-recall trade-off and business suitability.")
    ]
    for b, n in objs:
        p = tf3.add_paragraph()
        p.space_before = Pt(12)
        r1 = p.add_run()
        r1.text = "• " + b
        r1.font.bold = True
        r1.font.size = Pt(14)
        r1.font.color.rgb = ACCENT_BLUE
        r2 = p.add_run()
        r2.text = n
        r2.font.size = Pt(14)
        r2.font.color.rgb = DARK_TEXT

    s3.notes_slide.notes_text_frame.text = (
        "Our five objectives guide our methodology: survey 10 papers, replicate Ounacer's baseline Isolation Forest, analyze its false-positive flaw, "
        "implement validation threshold tuning, and rigorously measure the resulting trade-offs on untouched test data."
    )

    # SLIDE 4: Existing Research & Selected Case Study
    s4 = prs.slides.add_slide(blank_layout)
    add_header(s4, "3. Literature Review & Selected Case Study")
    tb4 = s4.shapes.add_textbox(Inches(0.8), Inches(1.7), Inches(11.7), Inches(5.2))
    tf4 = tb4.text_frame
    tf4.word_wrap = True

    p = tf4.paragraphs[0]
    p.text = "10 Investigated Peer-Reviewed Studies:"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = NAVY

    p_sub = tf4.add_paragraph()
    p_sub.text = "Covered IoT botnets (N-BaIoT), network intrusion (Kitsune), visual manufacturing inspection (MVTec AD), medical imaging (AnoGAN), spacecraft telemetry (NASA LSTM-NDT), cloud logs (DeepLog), and seasonal KPIs (Donut, HTM)."
    p_sub.font.size = Pt(13)
    p_sub.font.color.rgb = DARK_TEXT
    p_sub.space_after = Pt(10)

    p_sel = tf4.add_paragraph()
    p_sel.text = "Selected Case Study for Implementation & Improvement:"
    p_sel.font.size = Pt(16)
    p_sel.font.bold = True
    p_sel.font.color.rgb = ACCENT_BLUE

    box_text = (
        "• Paper Title: Using Isolation Forest in anomaly detection: the case of credit card transactions\n"
        "• Authors & Year: Ounacer, S., El Arass, M.Y., Gherabi, N. and Bahaj, M. (2018)\n"
        "• Publication Source: IEEE 4th International Conference on Optimization and Applications (ICOA 2018)\n"
        "• Published Findings: Ounacer et al. demonstrated that Isolation Forest scales linearly O(n) and runs orders of magnitude faster than distance-based LOF on card transaction data.\n"
        "• Identified Limitation & Difference: Ounacer et al. evaluated Isolation Forest using default threshold settings without validation tuning. In our replication experiment, this default boundary produces 1,996 false-positive alarms. Our project keeps their model fixed and optimizes the threshold on validation data to improve precision."
    )
    p_box = tf4.add_paragraph()
    p_box.text = box_text
    p_box.font.size = Pt(12.5)
    p_box.font.color.rgb = DARK_TEXT

    s4.notes_slide.notes_text_frame.text = (
        "We surveyed 10 papers and chose Ounacer et al. (2018). It is important to distinguish published findings from our own experiment: "
        "Ounacer et al. proved that Isolation Forest is exceptionally fast and scalable compared to LOF. "
        "However, they used default threshold settings. Our experimental replication revealed that this default boundary produces 1,996 false alarms. "
        "Our improvement specifically solves this thresholding limitation."
    )

    # SLIDE 5: Dataset & Class Imbalance
    s5 = prs.slides.add_slide(blank_layout)
    add_header(s5, "4. Dataset Architecture & Verified Properties")

    tb5_left = s5.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(5.8), Inches(5.0))
    tf5_l = tb5_left.text_frame
    tf5_l.word_wrap = True

    p = tf5_l.paragraphs[0]
    p.text = "European Credit Card Fraud Benchmark (ULB):"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = NAVY

    ds_points = [
        "Total Transactions: 284,807 transactions (over 48 hours in Sept 2013).",
        "Genuine Class (0): 284,315 transactions (99.828%).",
        "Fraudulent Class (1): 492 transactions (0.172%).",
        "Class Skew: Approximately 1 fraud per 578 genuine payments.",
        "Features: 30 continuous predictors (Time, Amount, and V1-V28 PCA components).",
        "Missing Values: Exactly 0 missing values.",
        "Duplicate Records: 1,081 duplicate rows. Because features are anonymized PCA components without account IDs, their exact cause cannot be verified from data alone; following standard practice, all rows are retained."
    ]
    for pt in ds_points:
        p = tf5_l.add_paragraph()
        p.space_before = Pt(6)
        p.text = "• " + pt
        p.font.size = Pt(12)
        p.font.color.rgb = DARK_TEXT

    img_path = "IBM3201_Group5_CreditCardFraud_Final/results/class_distribution.png"
    if os.path.exists(img_path):
        s5.shapes.add_picture(img_path, Inches(6.8), Inches(1.9), width=Inches(5.8))

    s5.notes_slide.notes_text_frame.text = (
        "Here are the verified dataset properties. Notice that the positive fraud class represents only 0.172% of transactions. "
        "We also verified that 1,081 duplicate rows exist. Rather than speculating on their cause, we note that anonymized PCA features "
        "prevent us from verifying whether they are re-attempts or system logs, so following standard academic practice, all rows are retained."
    )

    # SLIDE 6: Original Isolation Forest Method
    s6 = prs.slides.add_slide(blank_layout)
    add_header(s6, "5. Original Approach: Baseline Isolation Forest")
    tb6 = s6.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
    tf6 = tb6.text_frame
    tf6.word_wrap = True

    p = tf6.paragraphs[0]
    p.text = "Principle of Isolation Forest (Liu et al., 2008):"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = NAVY

    p2 = tf6.add_paragraph()
    p2.text = (
        "Unlike distance- or density-based anomaly detectors, Isolation Forest isolates anomalies explicitly using recursive random binary partitioning. "
        "Because anomalies are few and divergent, they require fewer random splits to isolate and have shorter average path lengths h(x)."
    )
    p2.font.size = Pt(13)
    p2.font.color.rgb = DARK_TEXT
    p2.space_before = Pt(6)

    algo_steps = [
        ("Ensemble Structure: ", "Constructs an ensemble of 100 iTrees, each built on a subsample of 256 instances."),
        ("Normalized Anomaly Score: ", "s(x, psi) = 2^(-E(h(x)) / c(psi)), where c(psi) is the average path length of unsuccessful searches in a Binary Search Tree."),
        ("Baseline Decision Boundary: ", "scikit-learn applies offset_ = -0.50, predicting fraud when suspicion score (-score_samples) >= 0.500."),
        ("Measured Baseline Performance on Test Set: ", "Catches 83 out of 98 fraud cases (Recall = 84.69%), but generates 1,996 false-positive alarms (Total alerts = 2,079). Precision is only 3.99% (96% false alarm rate).")
    ]
    for bold_txt, norm_txt in algo_steps:
        p = tf6.add_paragraph()
        p.space_before = Pt(8)
        run1 = p.add_run()
        run1.text = "• " + bold_txt
        run1.font.bold = True
        run1.font.size = Pt(13)
        run1.font.color.rgb = NAVY
        run2 = p.add_run()
        run2.text = norm_txt
        run2.font.size = Pt(13)
        run2.font.color.rgb = DARK_TEXT

    s6.notes_slide.notes_text_frame.text = (
        "Isolation Forest works by recursive random splitting. Anomalies isolate quickly near the root. "
        "However, when evaluated with the default 0.500 decision threshold, it catches 83 frauds but flags 1,996 innocent transactions as fraud! "
        "This gives a precision of only 3.99%, meaning 96% of all alerts are false alarms."
    )

    # SLIDE 7: Proposed Improvement: Threshold Tuning
    s7 = prs.slides.add_slide(blank_layout)
    add_header(s7, "6. Proposed Improvement: Validation Threshold Tuning")

    tb7_l = s7.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(6.0), Inches(5.0))
    tf7_l = tb7_l.text_frame
    tf7_l.word_wrap = True

    p = tf7_l.paragraphs[0]
    p.text = "Validation Tuning Methodology:"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = NAVY

    t_steps = [
        ("Frozen Isolation Forest Model: ", "The trained tree structures are kept fixed so the experiment isolates the exact effect of threshold tuning."),
        ("Continuous Suspicion Score: ", "Mapped as Suspicion(x) = -score_samples(x). Larger values denote higher anomaly likelihood."),
        ("Validation Grid Search: ", "Evaluates candidate thresholds across 1,000 percentiles (85th to 99.99th) strictly on the Validation set."),
        ("Optimization Objective: ", "Selects threshold tau* that maximizes validation F1-score: tau* = argmax F1(val, tau)."),
        ("Strict Leakage Protection: ", "The untouched test set is never seen during threshold selection. Optimal threshold selected is tau* = 0.6250.")
    ]
    for b, n in t_steps:
        p = tf7_l.add_paragraph()
        p.space_before = Pt(8)
        r1 = p.add_run()
        r1.text = "• " + b
        r1.font.bold = True
        r1.font.size = Pt(13)
        r1.font.color.rgb = GREEN
        r2 = p.add_run()
        r2.text = n
        r2.font.size = Pt(13)
        r2.font.color.rgb = DARK_TEXT

    img_path = "IBM3201_Group5_CreditCardFraud_Final/results/threshold_analysis.png"
    if os.path.exists(img_path):
        s7.shapes.add_picture(img_path, Inches(7.0), Inches(1.9), width=Inches(5.7))

    s7.notes_slide.notes_text_frame.text = (
        "Our improvement freezes the trained model and tunes the decision threshold on the independent validation set. "
        "As seen in the trajectory plot, increasing the threshold drives precision up while recall gently declines. "
        "The peak validation F1 occurs at 0.6250, shifting our decision boundary to heavily filter out false positives."
    )

    # SLIDE 8: Experimental Setup
    s8 = prs.slides.add_slide(blank_layout)
    add_header(s8, "7. Experimental Setup & Preprocessing Workflow")
    tb8 = s8.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
    tf8 = tb8.text_frame
    tf8.word_wrap = True

    p = tf8.paragraphs[0]
    p.text = "Rigorous 3-Way Stratified Partitioning:"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = NAVY

    wf_steps = [
        ("Stratified Split (60% / 20% / 20%): ", "Train (170,883 samples, 295 frauds) | Validation (56,962 samples, 99 frauds) | Test (56,962 samples, 98 frauds)."),
        ("Preprocessing Safeguards: ", "RobustScaler is fitted strictly on genuine training records (y_train == 0), preventing outlier distortion and test data leakage."),
        ("Training Constraint: ", "Isolation Forest is trained exclusively on unlabelled genuine transactions, mirroring real-world unsupervised/semi-supervised deployment."),
        ("Identical Test Records: ", "Both baseline and improved models are evaluated on the exact same 56,962 test transactions."),
        ("Verification & Reproducibility: ", "Fixed random seed (42); 8 automated unit tests pass in Pytest.")
    ]
    for b, n in wf_steps:
        p = tf8.add_paragraph()
        p.space_before = Pt(10)
        r1 = p.add_run()
        r1.text = "• " + b
        r1.font.bold = True
        r1.font.size = Pt(13)
        r1.font.color.rgb = ACCENT_BLUE
        r2 = p.add_run()
        r2.text = n
        r2.font.size = Pt(13)
        r2.font.color.rgb = DARK_TEXT

    s8.notes_slide.notes_text_frame.text = (
        "We followed strict experimental rigor. The data was split 60/20/20 with stratification. "
        "The scaler and Isolation Forest were fit exclusively on genuine training records. "
        "Both models were tested on the exact same 56,962 test records, guaranteeing a fair and reproducible comparison."
    )

    # SLIDE 9: Results & Performance Comparison Table
    s9 = prs.slides.add_slide(blank_layout)
    add_header(s9, "8. Experimental Results & Performance Comparison")

    rows, cols = 10, 3
    left, top, width, height = Inches(0.8), Inches(1.7), Inches(11.7), Inches(5.0)
    table_shape = s9.shapes.add_table(rows, cols, left, top, width, height)
    tbl = table_shape.table

    tbl.columns[0].width = Inches(4.5)
    tbl.columns[1].width = Inches(3.6)
    tbl.columns[2].width = Inches(3.6)

    headers = ["Evaluation Metric", "Original Baseline (Threshold = 0.500)", "Threshold-Tuned (Threshold = 0.625)"]
    for j, h in enumerate(headers):
        cell = tbl.cell(0, j)
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY
        p = cell.text_frame.paragraphs[0]
        p.text = h
        p.font.bold = True
        p.font.size = Pt(13)
        p.font.color.rgb = WHITE
        p.alignment = PP_ALIGN.CENTER if j > 0 else PP_ALIGN.LEFT

    table_data = [
        ("Precision", "0.0399 (3.99%)", "0.2174 (21.74%) [+444.5%]"),
        ("Recall", "0.8469 (84.69%)", "0.4082 (40.82%) [-51.8%]"),
        ("F1-Score", "0.0763", "0.2837 [+271.8%]"),
        ("Accuracy", "0.9647 (96.47%)", "0.9965 (99.65%)"),
        ("False Positive Rate (FPR)", "0.035101 (3.51%)", "0.002532 (0.25%) [-92.79%]"),
        ("False Negative Rate (FNR)", "0.1531 (15.31%)", "0.5918 (59.18%)"),
        ("Average Precision / PR-AUC", "0.1715", "0.1715"),
        ("False-Positive Alerts (FP)", "1,996 alerts", "144 alerts [1,852 fewer FPs, -92.79%]"),
        ("Total Alerts Generated (TP + FP)", "2,079 alerts", "184 alerts [1,895 fewer alerts, -91.15%]")
    ]

    for i, (m, orig, tuned) in enumerate(table_data):
        row_idx = i + 1
        c0 = tbl.cell(row_idx, 0)
        c1 = tbl.cell(row_idx, 1)
        c2 = tbl.cell(row_idx, 2)
        c0.text = m
        c1.text = orig
        c2.text = tuned
        for cell, j in zip([c0, c1, c2], [0, 1, 2]):
            p = cell.text_frame.paragraphs[0]
            p.font.size = Pt(12)
            if j == 2 and ("+" in tuned or "fewer" in tuned):
                p.font.bold = True
                p.font.color.rgb = GREEN if ("+" in tuned or "fewer" in tuned) else DARK_TEXT
            if j > 0:
                p.alignment = PP_ALIGN.CENTER

    s9.notes_slide.notes_text_frame.text = (
        "Here are our actual calculated results. Notice the critical distinctions: "
        "Total alerts generated were reduced by 1,895 (a 91.15% reduction from 2,079 to 184). "
        "False-positive alerts were reduced by 1,852 (a 92.79% reduction from 1,996 to 144). "
        "Precision increased from 3.99% to 21.74%, but recall dropped from 84.69% to 40.82%."
    )

    # SLIDE 10: Visual Analysis & Confusion Matrices
    s10 = prs.slides.add_slide(blank_layout)
    add_header(s10, "9. Confusion Matrices & Precision-Recall Curve")

    cm_orig_img = "IBM3201_Group5_CreditCardFraud_Final/results/confusion_matrix_original.png"
    cm_tuned_img = "IBM3201_Group5_CreditCardFraud_Final/results/confusion_matrix_improved.png"
    pr_curve_img = "IBM3201_Group5_CreditCardFraud_Final/results/precision_recall_curve.png"

    if os.path.exists(cm_orig_img) and os.path.exists(cm_tuned_img):
        s10.shapes.add_picture(cm_orig_img, Inches(0.8), Inches(1.8), width=Inches(3.8))
        s10.shapes.add_picture(cm_tuned_img, Inches(4.8), Inches(1.8), width=Inches(3.8))
    if os.path.exists(pr_curve_img):
        s10.shapes.add_picture(pr_curve_img, Inches(8.8), Inches(1.8), width=Inches(4.1))

    s10.notes_slide.notes_text_frame.text = (
        "The confusion matrices clearly display the operational shift: "
        "The baseline caught 83 frauds but triggered 1,996 false alarms. "
        "The tuned model slashed false alarms to 144, but caught 40 frauds (missing 58). "
        "On the right, the Precision-Recall curve shows how threshold tuning shifts our operating point along the trade-off curve."
    )

    # SLIDE 11: Honest Trade-off & Business Suitability
    s11 = prs.slides.add_slide(blank_layout)
    add_header(s11, "10. Honest Trade-off Analysis & Business Suitability")
    tb11 = s11.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
    tf11 = tb11.text_frame
    tf11.word_wrap = True

    p = tf11.paragraphs[0]
    p.text = "Operational Realities & Decision Context:"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = NAVY

    disc_points = [
        ("The Inescapable Trade-Off: ", "Lowering false alarms reduced recall from 84.69% to 40.82%. The tuned model misses 43 more fraud cases than the baseline (58 missed vs. 15 missed)."),
        ("Why Higher F1 is Not Universally Superior: ", "An F1 of 0.2837 is mathematically higher than 0.0763, but does not fit every business model. The best threshold depends on the financial cost of missed fraud versus the cost of investigating false alarms."),
        ("When Baseline (Lower Threshold) is Preferred: ", "High-value commercial transactions, wire transfers, or wealth banking where a single missed fraud costs millions. Banks will pay investigators to check 2,000 alerts to catch 83 frauds."),
        ("When Tuned (Higher Threshold) is Preferred: ", "High-volume retail debit cards where fraud review teams have strict daily limits and false declines frustrate customers. Slashing false alarms by 92.79% makes the alert queue manageable."),
        ("The Accuracy Paradox: ", "Both models show >96% accuracy because genuine transactions dominate (99.828%). Predicting all records as genuine yields 99.83% accuracy while catching zero fraud.")
    ]
    for b, n in disc_points:
        p = tf11.add_paragraph()
        p.space_before = Pt(8)
        r1 = p.add_run()
        r1.text = "• " + b
        r1.font.bold = True
        r1.font.size = Pt(13)
        r1.font.color.rgb = RED if "Trade-Off" in b or "Missed" in b else NAVY
        r2 = p.add_run()
        r2.text = n
        r2.font.size = Pt(13)
        r2.font.color.rgb = DARK_TEXT

    s11.notes_slide.notes_text_frame.text = (
        "We must be honest about limitations: tuning increased precision by dropping recall from 84.7% to 40.8%. "
        "The tuned model misses 58 frauds, which is 43 more missed frauds than the baseline. "
        "Therefore, a higher F1 does not make it universally better for every bank. "
        "Banks must select thresholds based on their investigation capacity and the cost of missed fraud."
    )

    # SLIDE 12: Conclusion & References
    s12 = prs.slides.add_slide(blank_layout)
    add_header(s12, "11. Conclusion & Verified References")
    tb12 = s12.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
    tf12 = tb12.text_frame
    tf12.word_wrap = True

    p = tf12.paragraphs[0]
    p.text = "Summary of Project Findings:"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = NAVY

    conc_points = [
        ("Empirical Replication: ", "Replicated Ounacer et al. (2018) and verified that default Isolation Forest thresholds produce massive false alarms (1,996 FPs) on imbalanced credit card data."),
        ("Demonstrated Gain: ", "Validation threshold tuning (0.6250) boosted precision by 444.5% and reduced false positives by 92.79% with zero test-set data leakage."),
        ("Engineering Takeaway: ", "Threshold tuning is a zero-latency, highly effective intervention that allows banking operations to calibrate false alarm rates against fraud loss risk.")
    ]
    for b, n in conc_points:
        p = tf12.add_paragraph()
        p.space_before = Pt(6)
        r1 = p.add_run()
        r1.text = "• " + b
        r1.font.bold = True
        r1.font.size = Pt(13)
        r1.font.color.rgb = GREEN
        r2 = p.add_run()
        r2.text = n
        r2.font.size = Pt(13)
        r2.font.color.rgb = DARK_TEXT

    p_ref = tf12.add_paragraph()
    p_ref.text = "\nKey Harvard References:"
    p_ref.font.size = Pt(14)
    p_ref.font.bold = True
    p_ref.font.color.rgb = ACCENT_BLUE

    refs = [
        "• Ounacer, S. et al. (2018) 'Using Isolation Forest in anomaly detection: the case of credit card transactions', IEEE ICOA, pp. 1-6.",
        "• Liu, F.T., Ting, K.M. and Zhou, Z.-H. (2008) 'Isolation Forest', 8th IEEE ICDM, pp. 413-422.",
        "• Carcillo, F. et al. (2021) 'Combining unsupervised and supervised learning in credit card fraud detection', Information Sciences, 557, pp. 317-331.",
        "• Saito, T. and Rehmsmeier, M. (2015) 'The Precision-Recall Plot Is More Informative than the ROC Plot...', PLOS ONE, 10(3), p. e0118432."
    ]
    for r in refs:
        p = tf12.add_paragraph()
        p.text = r
        p.font.size = Pt(11)
        p.font.color.rgb = DARK_TEXT

    s12.notes_slide.notes_text_frame.text = (
        "In conclusion, our project demonstrates that threshold tuning is a zero-cost operational intervention that empowers financial institutions "
        "to navigate the precision-recall trade-off intentionally. Thank you, and we welcome your questions!"
    )

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    prs.save(output_path)
    print(f"PowerPoint presentation created successfully at: {output_path}")

if __name__ == '__main__':
    create_deck('IBM3201_Group5_CreditCardFraud_Final/presentation.pptx')
